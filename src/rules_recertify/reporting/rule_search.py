"""Search the latest ingested rule snapshot from a user-supplied item list."""
from __future__ import annotations

import csv
import importlib.util
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from ..history.database import Database
from ..workloader.csvio import read_rows
from .workbook import _expand_services, _load_report_services, _raw_reference_candidates


SEARCH_FIELDS = (
    "ruleset_scope", "rule_type", "src_all_workloads",
    "src_labels", "src_labels_exclusions", "src_label_groups",
    "src_label_groups_exclusions", "src_iplists", "src_workloads",
    "dst_all_workloads", "dst_labels", "dst_labels_exclusions", "dst_label_groups",
    "dst_label_groups_exclusions", "dst_iplists", "dst_workloads", "services",
)

_PORT_TOKEN = re.compile(
    r"^\s*(?:(TCP|UDP)\s*/\s*)?(\d{1,5})(?:\s*-\s*(\d{1,5}))?"
    r"(?:\s+(TCP|UDP))?\s*$", re.IGNORECASE,
)
_RULE_PORT = re.compile(
    r"\b(\d{1,5})(?:\s*-\s*(\d{1,5}))?\s+(TCP|UDP)\b", re.IGNORECASE,
)
LABEL_GROUP_FIELDS = (
    "name", "key", "description", "member_labels", "member_label_groups",
    "fully_expanded_members", "href",
)
LABEL_TYPE_NAMES = {
    "app": "label:app",
    "env": "label:env",
    "role": "label:role",
    "loc": "label:location",
    "location": "label:location",
    "os": "label:OS",
    "module": "label:module",
}
SHEET_HEADER_COLORS = {
    "Summary": "1F4E78",
    "Results": "305496",
    "label-groups": "548235",
    "Metadata": "5B5B5B",
    "Bouquets": "1F4E78",
    "iplists": "C65911",
    "workloads": "548235",
    "labelgroups": "7030A0",
}
RESULT_RULE_FIELDS = (
    "ruleset_name", "ruleset_description", "ruleset_scope", "ruleset_enabled",
    "rule_type", "rule_description", "rule_enabled", "unscoped_consumers",
    "src_all_workloads", "src_labels", "src_labels_exclusions", "src_label_groups",
    "src_label_groups_exclusions", "src_iplists", "src_user_groups", "src_workloads",
    "src_virtual_services", "src_use_workload_subnets", "dst_all_workloads", "dst_labels",
    "dst_labels_exclusions", "dst_label_groups", "dst_label_groups_exclusions", "dst_iplists",
    "dst_user_groups", "dst_virtual_services", "dst_workloads", "dst_use_workload_subnets",
    "services", "src_resolve_labels_as", "dst_resolve_labels_as", "machine_auth_enabled",
    "secure_connect_enabled", "stateless_enabled", "network_type", "external_data_set",
    "external_data_reference", "ruleset_href", "rule_href", "update_type",
)


def _contains_whole_item(value: object, item: str, case_sensitive: bool) -> bool:
    flags = 0 if case_sensitive else re.IGNORECASE
    pattern = rf"(?<![A-Za-z0-9_]){re.escape(item)}(?![A-Za-z0-9_])"
    return re.search(pattern, str(value or ""), flags) is not None


def load_search_items(path: Path) -> List[str]:
    """Load one non-empty item per CSV/TXT row or XLSX row."""
    if not path.is_file():
        raise ValueError(f"items file not found: {path}")
    if path.suffix.lower() == ".xlsx":
        if importlib.util.find_spec("openpyxl") is None:
            raise RuntimeError("openpyxl is required to read an XLSX items file")
        from openpyxl import load_workbook
        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = [[cell for cell in row if cell is not None and str(cell).strip()] for row in sheet.iter_rows(values_only=True)]
    else:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
    items: List[str] = []
    for number, row in enumerate(rows, 1):
        values = [str(value).strip() for value in row if str(value).strip()]
        if not values:
            continue
        if len(values) != 1:
            raise ValueError(f"{path}:{number}: expected exactly one non-empty column")
        if number == 1 and values[0].casefold() in {"item", "search item", "searched item"}:
            continue
        if values[0] not in items:
            items.append(values[0])
    if not items:
        raise ValueError("items file is empty or contains only blank values")
    return items


def latest_rules(db: Database) -> Tuple[List[Dict[str, object]], str]:
    """Return only rules explicitly present in the current policy inventory."""
    with db.connect() as connection:
        latest = connection.execute(
            "SELECT MAX(snapshot_at) FROM rules WHERE is_present=1"
        ).fetchone()[0]
        if not latest:
            return [], ""
        rows = [dict(row) for row in connection.execute(
            "SELECT * FROM rules WHERE is_present=1 ORDER BY ruleset_name,rule_href",
        )]
    return rows, str(latest)


def _port_query(item: str) -> Optional[List[Tuple[Optional[str], int, int]]]:
    tokens = [token.strip() for token in re.split(r"[;,]", item) if token.strip()]
    parsed: List[Tuple[Optional[str], int, int]] = []
    if not tokens:
        return None
    for token in tokens:
        match = _PORT_TOKEN.fullmatch(token)
        if not match:
            return None
        protocol = (match.group(1) or match.group(4) or "").upper() or None
        start, end = int(match.group(2)), int(match.group(3) or match.group(2))
        if start > end or end > 65535:
            raise ValueError(f"invalid port or range: {token}")
        parsed.append((protocol, start, end))
    return parsed


def _service_intersections(
    expanded_services: str, query: Sequence[Tuple[Optional[str], int, int]],
) -> List[str]:
    rule_intervals: List[Tuple[str, int, int]] = []
    for match in _RULE_PORT.finditer(expanded_services):
        start, end = int(match.group(1)), int(match.group(2) or match.group(1))
        if start <= end <= 65535:
            rule_intervals.append((match.group(3).upper(), start, end))
    matches: List[str] = []
    for query_protocol, query_start, query_end in query:
        for rule_protocol, rule_start, rule_end in rule_intervals:
            if query_protocol and query_protocol != rule_protocol:
                continue
            start, end = max(query_start, rule_start), min(query_end, rule_end)
            if start <= end:
                value = f"{rule_protocol}/{start}" if start == end else f"{rule_protocol}/{start}-{end}"
                if value not in matches:
                    matches.append(value)
    return matches


def search_latest_rules(
    rules: Iterable[Mapping[str, object]], items: Sequence[str],
    service_catalog: Mapping[str, Tuple[str, str]], case_sensitive: bool = False,
) -> Tuple[List[Dict[str, object]], List[Dict[str, object]]]:
    """Return result and summary rows, one result per searched item/rule."""
    prepared = []
    for rule in rules:
        raw = json.loads(str(rule["raw_json"]))
        prepared.append((rule, raw))
    results: List[Dict[str, object]] = []
    summary: List[Dict[str, object]] = []
    for item in items:
        port_query = _port_query(item)
        matched_count = 0
        for rule, raw in prepared:
            found_fields: List[str] = []
            details: List[str] = []
            if port_query is not None:
                expanded = _expand_services(str(raw.get("services", rule.get("services", ""))), service_catalog)
                intersections = _service_intersections(expanded, port_query)
                if intersections:
                    found_fields.append("services")
                    details.extend(intersections)
            else:
                for field in SEARCH_FIELDS:
                    value = str(raw.get(field, rule.get(field, "")) or "")
                    if _contains_whole_item(value, item, case_sensitive):
                        found_fields.append(field)
            if found_fields:
                matched_count += 1
                results.append({
                    "searched item": item,
                    "status": "FOUND",
                    "found as": ";".join(found_fields),
                    "match details": ";".join(details),
                    "rule snapshot": rule["snapshot_at"],
                    **{header: raw.get(header, rule.get(header, "")) for header in RESULT_RULE_FIELDS},
                })
        if not matched_count:
            results.append({
                "searched item": item, "status": "NOT_USED_IN_ANY_RULE",
                "found as": "", "match details": "", "rule snapshot": "",
                **{header: "" for header in RESULT_RULE_FIELDS},
            })
        summary.append({
            "searched item": item,
            "status": "FOUND" if matched_count else "NOT_USED_IN_ANY_RULE",
            "matching rules": matched_count,
        })
    results.sort(key=lambda row: (0 if row["status"] == "NOT_USED_IN_ANY_RULE" else 1,
                                  str(row["searched item"]).casefold(), str(row.get("rule_href", ""))))
    return results, summary


def _load_label_groups(raw_dir: Optional[Path]) -> Tuple[List[Dict[str, str]], str]:
    """Load label groups from the materialized snapshot, then recent raw runs."""
    if not raw_dir:
        return [], ""
    candidates = _raw_reference_candidates(raw_dir, "export_labelgroups.csv")
    if not candidates:
        return [], ""
    source = candidates[0]
    return list(read_rows(source, LABEL_GROUP_FIELDS)), str(source)


def classify_search_items(
    items: Sequence[str], raw_dir: Optional[Path],
    label_groups: Sequence[Mapping[str, str]], case_sensitive: bool = False,
) -> Dict[str, str]:
    """Classify exact item names from the current snapshot reference exports."""
    types: Dict[str, List[str]] = {item: [] for item in items}

    def equals(left: object, right: str) -> bool:
        value = str(left or "").strip()
        return value == right if case_sensitive else value.casefold() == right.casefold()

    def add(item: str, value: str) -> None:
        if value and value not in types[item]:
            types[item].append(value)

    if raw_dir:
        ip_list_candidates = _raw_reference_candidates(raw_dir, "export_iplists.csv")
        if ip_list_candidates:
            for row in read_rows(ip_list_candidates[0], ("name", "include")):
                for item in items:
                    if equals(row["name"], item):
                        add(item, "iplist")

        label_candidates = _raw_reference_candidates(raw_dir, "labels.csv")
        if label_candidates:
            for row in read_rows(label_candidates[0], ("key", "value")):
                label_type = LABEL_TYPE_NAMES.get(row["key"].strip().casefold())
                if not label_type:
                    continue
                for item in items:
                    if equals(row["value"], item):
                        add(item, label_type)

    for group in label_groups:
        for item in items:
            if equals(group.get("name", ""), item):
                add(item, "labelgroup")

    return {item: ";".join(values) for item, values in types.items()}


def search_label_groups(
    label_groups: Sequence[Mapping[str, str]], items: Sequence[str],
    case_sensitive: bool = False,
) -> List[Dict[str, str]]:
    """Return direct and transitive label-group membership for each item."""
    output: List[Dict[str, str]] = []

    def contains(value: object, needle: str) -> bool:
        return _contains_whole_item(value, needle, case_sensitive)

    for item in items:
        matched: set[int] = set()
        found_as: Dict[int, List[str]] = {}
        for index, group in enumerate(label_groups):
            fields = [
                field for field in ("member_labels", "fully_expanded_members")
                if contains(group.get(field, ""), item)
            ]
            if fields:
                matched.add(index)
                found_as[index] = fields

        # Some Workloader versions do not fully expand nested groups. Walk the
        # parent references as a fallback, using either the child name or href.
        changed = True
        while changed:
            changed = False
            children = [label_groups[index] for index in matched]
            for index, group in enumerate(label_groups):
                if index in matched:
                    continue
                memberships = group.get("member_label_groups", "")
                if any(
                    (child.get("name") and contains(memberships, child["name"]))
                    or (child.get("href") and contains(memberships, child["href"]))
                    for child in children
                ):
                    matched.add(index)
                    found_as[index] = ["member_label_groups (transitive)"]
                    changed = True

        if not matched:
            output.append({
                "searched item": item,
                "status": "NOT_USED_IN_ANY_LABEL_GROUP",
                "found as": "",
                **{field: "" for field in LABEL_GROUP_FIELDS},
            })
            continue
        for index in sorted(matched, key=lambda value: (
            str(label_groups[value].get("name", "")).casefold(),
            str(label_groups[value].get("href", "")),
        )):
            group = label_groups[index]
            output.append({
                "searched item": item,
                "status": "FOUND",
                "found as": ";".join(found_as[index]),
                **{field: str(group.get(field, "") or "") for field in LABEL_GROUP_FIELDS},
            })
    return output


def _write_sheet(workbook: object, name: str, rows: Sequence[Mapping[str, object]]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    sheet = workbook.create_sheet(name)
    headers = list(rows[0]) if rows else ["searched item", "status"]
    sheet.append(headers)
    for row in rows:
        sheet.append([_excel_value(row.get(header, "")) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill(
        fill_type="solid", fgColor=SHEET_HEADER_COLORS.get(name, "305496"),
    )
    header_border = Border(bottom=Side(style="thin", color="FFFFFF"))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
    sheet.row_dimensions[1].height = 24

    if "status" in headers:
        status_column = headers.index("status") + 1
        for row_number in range(2, sheet.max_row + 1):
            cell = sheet.cell(row_number, status_column)
            if cell.value == "FOUND":
                cell.fill = PatternFill(fill_type="solid", fgColor="C6E0B4")
            elif str(cell.value or "").startswith("NOT_USED"):
                cell.fill = PatternFill(fill_type="solid", fgColor="F4CCCC")
            cell.font = Font(bold=True)

    for index, header in enumerate(headers, 1):
        lengths = [len(str(header))] + [len(str(row.get(header, "") or "")) for row in rows]
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(lengths) + 2, 80)


def _excel_value(value: object) -> object:
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, sort_keys=True, ensure_ascii=False)
    return value


def generate_rule_search_workbook(
    db: Database, items_path: Path, output: Path, raw_dir: Optional[Path] = None,
    case_sensitive: bool = False,
) -> Path:
    if importlib.util.find_spec("openpyxl") is None:
        raise RuntimeError("openpyxl is required for rule-search Excel generation")
    from openpyxl import Workbook
    items = load_search_items(items_path)
    rules, snapshot = latest_rules(db)
    if not snapshot:
        raise ValueError("no rule snapshot is available in SQLite")
    service_catalog, service_reference = _load_report_services(raw_dir)
    label_groups, label_group_reference = _load_label_groups(raw_dir)
    results, summary = search_latest_rules(rules, items, service_catalog, case_sensitive)
    item_types = classify_search_items(items, raw_dir, label_groups, case_sensitive)
    summary = [
        {
            "searched item": row["searched item"],
            "type": item_types.get(str(row["searched item"]), ""),
            "status": row["status"],
            "matching rules": row["matching rules"],
        }
        for row in summary
    ]
    label_group_results = search_label_groups(label_groups, items, case_sensitive)
    workbook = Workbook(); workbook.remove(workbook.active)
    _write_sheet(workbook, "Summary", summary)
    _write_sheet(workbook, "Results", results)
    _write_sheet(workbook, "label-groups", label_group_results)
    _write_sheet(workbook, "Metadata", [{
        "rule snapshot": snapshot,
        "items file": str(items_path),
        "service reference": service_reference,
        "label-group reference": label_group_reference or "NOT_AVAILABLE",
        "case sensitive": str(case_sensitive).upper(),
        "generated at UTC": datetime.now(timezone.utc).isoformat(),
    }])
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    workbook.save(temporary); temporary.replace(output)
    return output
