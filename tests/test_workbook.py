import json
import importlib.util
import csv
import tempfile
import unittest
from datetime import date
from pathlib import Path

from rules_recertify.history.database import Database
from rules_recertify.reporting.workbook import (
    _count_addresses, _count_ports, _excel_safe_count, _expand_side,
    _expand_services, _expand_side_details, _load_report_ip_lists, _rule_matches, _scope_pairs,
    generate_workbook,
)


class WorkbookExpansionTest(unittest.TestCase):
    def setUp(self):
        self.workloads = [
            {
                "hostname": "payment-01.example.net",
                "short_hostname": "PAYMENT-01",
                "name": "payment-01",
                "app": "APM_PAYMENT",
                "env": "PRD",
                "addresses_json": json.dumps(["10.10.1.10"]),
            },
            {
                "hostname": "payment-02.example.net",
                "short_hostname": "PAYMENT-02",
                "name": "payment-02",
                "app": "APM_PAYMENT",
                "env": "PRD",
                "addresses_json": json.dumps(["10.10.1.11"]),
            },
            {
                "hostname": "payment-uat.example.net",
                "short_hostname": "PAYMENT-UAT",
                "name": "payment-uat",
                "app": "APM_PAYMENT",
                "env": "UAT",
                "addresses_json": json.dumps(["10.20.1.10"]),
            },
            {
                "hostname": "other.example.net",
                "short_hostname": "OTHER",
                "name": "other",
                "app": "OTHER_APP",
                "env": "PRD",
                "addresses_json": json.dumps(["10.30.1.10"]),
            },
        ]

    def test_all_workloads_expands_app_and_environment_scope(self):
        expanded = _expand_side(
            {
                "ruleset_scope": "app:APM_PAYMENT;env:PRD",
                "src_all_workloads": "true",
            },
            "src",
            self.workloads,
            "PRD",
        )
        self.assertEqual(
            expanded.splitlines(),
            [
                "PAYMENT-01 (10.10.1.10)",
                "PAYMENT-02 (10.10.1.11)",
            ],
        )

    def test_all_workloads_uses_every_label_from_ruleset_scope(self):
        workloads = [
            {**self.workloads[0], "loc": "PAR", "role": "WEB"},
            {**self.workloads[1], "loc": "LYO", "role": "WEB"},
        ]
        expanded, addresses = _expand_side_details(
            {
                "ruleset_scope": "app:APM_PAYMENT;env:PRD;loc:PAR;role:WEB",
                "src_all_workloads": "true",
            },
            "src", workloads, "PRD",
        )
        self.assertEqual(expanded, "PAYMENT-01 (10.10.1.10)")
        self.assertEqual(addresses, ["10.10.1.10"])

    def test_all_workloads_expansion_accepts_reversed_scope_order(self):
        expanded = _expand_side(
            {
                "ruleset_scope": "env:PRD;app:APM_PAYMENT",
                "dst_all_workloads": "TRUE",
            },
            "dst",
            self.workloads,
            "PRD",
        )

        self.assertNotIn("payment-uat", expanded)
        self.assertNotIn("other.example.net", expanded)
        self.assertIn("PAYMENT-01 (10.10.1.10)", expanded)

    def test_null_scope_environment_uses_requested_report_environment(self):
        expanded = _expand_side(
            {
                "ruleset_scope": "app:APM_PAYMENT;env:NULL",
                "src_all_workloads": "true",
            },
            "src",
            self.workloads,
            "UAT",
        )

        self.assertEqual(expanded, "PAYMENT-UAT (10.20.1.10)")

    def test_label_workload_uses_short_name_and_groups_multiple_addresses(self):
        workload = {
            "href": "/workloads/1", "hostname": "host.example.net",
            "short_hostname": "HOST", "name": "fallback", "app": "APP",
            "env": "PRD", "loc": "PAR", "role": "WEB",
            "addresses_json": json.dumps(["192.168.1.10", "192.168.1.5"]),
        }
        expanded = _expand_side(
            {"src_labels": "app:app;env:prd;role:web"}, "src", [workload], "prd"
        )
        self.assertEqual(expanded, "HOST (192.168.1.10;192.168.1.5)")

    def test_empty_short_hostname_falls_back_to_name(self):
        workload = {
            "href": "/workloads/1", "hostname": "", "short_hostname": "",
            "name": "fallback-name", "app": "APP", "env": "PRD",
            "addresses_json": json.dumps(["192.168.1.10"]),
        }
        expanded = _expand_side(
            {"dst_workloads": "/workloads/1"}, "dst", [workload], "PRD"
        )
        self.assertEqual(expanded, "fallback-name (192.168.1.10)")

    def test_ip_list_expands_members_and_removes_comments(self):
        expanded, addresses = _expand_side_details(
            {"src_iplists": "NETWORKS"}, "src", [], "PRD",
            [
                {"name": "NETWORKS", "member": "171.18.16.0/24#GEN1"},
                {"name": "NETWORKS", "member": "171.16.16.0/24"},
            ],
        )
        self.assertEqual(
            expanded,
            "IP List: NETWORKS (171.18.16.0/24;171.16.16.0/24)",
        )
        self.assertEqual(addresses, ["171.18.16.0/24", "171.16.16.0/24"])

    def test_unresolved_ip_list_is_visible(self):
        expanded = _expand_side(
            {"dst_iplists": "MISSING"}, "dst", [], "PRD", []
        )
        self.assertEqual(expanded, "IP List: MISSING [unresolved]")

    def test_ip_list_selector_accepts_workloader_display_and_href_annotation(self):
        expanded = _expand_side(
            {"src_iplists": "IP List: NETWORKS (/orgs/1/sec_policy/draft/ip_lists/1)"},
            "src", [], [("APP", "PRD")],
            [{"name": "NETWORKS", "include": "192.168.19.0/24#GEN1"}],
        )
        self.assertEqual(expanded, "IP List: NETWORKS (192.168.19.0/24)")

    def test_ip_list_include_splits_members_and_each_inline_comment(self):
        expanded = _expand_side(
            {"dst_iplists": "NETWORKS: /orgs/1/sec_policy/draft/ip_lists/1"},
            "dst", [], [("APP", "PRD")],
            [{"name": "NETWORKS", "include": "192.168.19.0/24#GEN1;192.168.19.1#GEN2"}],
        )
        self.assertEqual(expanded, "IP List: NETWORKS (192.168.19.0/24;192.168.19.1)")

    def test_report_prefers_latest_complete_raw_ip_list_export_over_stale_database(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); raw_dir = root / "raw"
            preflight = raw_dir / "preflight"; preflight.mkdir(parents=True)
            old_run = raw_dir / "20260910T000000Z-aabbccdd"; old_run.mkdir(parents=True)
            latest_run = raw_dir / "20260911T094511Z-331e4c5c"; latest_run.mkdir()
            (preflight / "export_iplists.csv").write_text(
                "name,include\nPREFLIGHT,203.0.113.0/24\n", encoding="utf-8",
            )
            with (old_run / "export_iplists.csv").open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=["name", "include"])
                writer.writeheader(); writer.writerow({"name": "OLD", "include": "10.0.0.0/8"})
            with (latest_run / "export_iplists.csv").open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=[
                    "name", "description", "include", "exclude", "ip_ranges",
                    "external_data_set", "external_data_ref", "href",
                ])
                writer.writeheader()
                writer.writerow({
                    "name": "NZ2_DOCKER.EUR.PRD_Paris.PCOM.L1-IPL",
                    "description": "infra vlan common name",
                    "include": "171.70.40.0/21#GEN1;171.71.32.0/21#GEN2",
                    "href": "/orgs/1/sec_policy/draft/ip_lists/1083035",
                })
            db = Database(root / "state.sqlite"); db.initialize()
            with db.connect() as connection:
                connection.execute("INSERT INTO ip_lists VALUES(?,?,?)", ("NZ3_ONLY", "192.0.2.0/24", "old"))
                rows, source = _load_report_ip_lists(connection, raw_dir)
            self.assertEqual(source, str(latest_run / "export_iplists.csv"))
            self.assertEqual(rows, [
                {"name": "NZ2_DOCKER.EUR.PRD_Paris.PCOM.L1-IPL", "member": "171.70.40.0/21"},
                {"name": "NZ2_DOCKER.EUR.PRD_Paris.PCOM.L1-IPL", "member": "171.71.32.0/21"},
            ])
            expanded = _expand_side(
                {"src_iplists": "NZ2_DOCKER.EUR.PRD_Paris.PCOM.L1-IPL"},
                "src", [], [("APP", "PRD")], rows,
            )
            self.assertEqual(
                expanded,
                "IP List: NZ2_DOCKER.EUR.PRD_Paris.PCOM.L1-IPL (171.70.40.0/21;171.71.32.0/21)",
            )

    def test_report_uses_sqlite_ip_lists_when_no_raw_export_exists(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); db = Database(root / "state.sqlite"); db.initialize()
            with db.connect() as connection:
                connection.execute("INSERT INTO ip_lists VALUES(?,?,?)", ("DATABASE_IPL", "192.0.2.0/24", "now"))
                rows, source = _load_report_ip_lists(connection, root / "missing-raw")
            self.assertEqual(rows[0]["name"], "DATABASE_IPL")
            self.assertEqual(source, "SQLite ip_lists fallback")

    def test_rule_selection_uses_exact_pairs_for_scoped_rulesets(self):
        pairs = [("APP_A", "PRD"), ("APP_B", "UAT")]
        matching = {"raw_json": json.dumps({"ruleset_scope": "app:APP_B;env:UAT"})}
        wrong_environment = {"raw_json": json.dumps({"ruleset_scope": "app:APP_B;env:PRD"})}
        self.assertTrue(_rule_matches(matching, pairs))
        self.assertFalse(_rule_matches(wrong_environment, pairs))

    def test_unscoped_rule_requires_pair_on_same_side_or_application_alone(self):
        pairs = [("APP_A", "PRD")]
        wrong_environment = {"raw_json": json.dumps({
            "ruleset_scope": "", "src_labels": "app:APP_A;env:UAT",
        })}
        split_across_sides = {"raw_json": json.dumps({
            "ruleset_scope": "", "src_labels": "app:APP_A", "dst_labels": "env:UAT",
        })}
        app_only = {"raw_json": json.dumps({"ruleset_scope": "", "dst_labels": "app:APP_A"})}
        self.assertFalse(_rule_matches(wrong_environment, pairs))
        self.assertTrue(_rule_matches(split_across_sides, pairs))
        self.assertTrue(_rule_matches(app_only, pairs))

    def test_multi_environment_expansion_does_not_cross_application_pairs(self):
        workloads = [
            {"short_hostname": "B-PRD", "name": "", "app": "APP_B", "env": "PRD",
             "addresses_json": json.dumps(["10.0.0.1"])},
            {"short_hostname": "B-UAT", "name": "", "app": "APP_B", "env": "UAT",
             "addresses_json": json.dumps(["10.0.0.2"])},
        ]
        expanded = _expand_side(
            {"src_labels": "app:APP_B"}, "src", workloads,
            [("APP_A", "PRD"), ("APP_B", "UAT")],
        )
        self.assertEqual(expanded, "B-UAT (10.0.0.2)")

    def test_scope_pair_cardinality_is_validated(self):
        with self.assertRaisesRegex(ValueError, "one-to-one"):
            _scope_pairs(["APP_A", "APP_B"], ["PRD"])

    def test_address_counts_are_stably_deduplicated(self):
        workload = {
            "href": "/workloads/1", "short_hostname": "HOST", "name": "host",
            "app": "APP", "env": "PRD", "addresses_json": json.dumps(["10.0.0.1", "10.0.0.2"]),
        }
        _, addresses = _expand_side_details(
            {"src_workloads": "/workloads/1", "src_labels": "app:APP;env:PRD"},
            "src", [workload], "PRD",
        )
        self.assertEqual(addresses, ["10.0.0.1", "10.0.0.2"])

    def test_port_count_expands_ranges_and_deduplicates_protocol_ports(self):
        self.assertEqual(_count_ports("443 TCP; 80-82 TCP; 53 UDP; 443 TCP"), 5)
        self.assertEqual(_count_ports("All Services; 0 ICMP"), 131072)
        self.assertEqual(_expand_services("All Services"), "0-65535 TCP;0-65535 UDP")

    def test_address_count_uses_network_cardinality_and_removes_overlap(self):
        self.assertEqual(_count_addresses(["192.168.19.0/24", "192.168.19.1", "10.0.0.1-10.0.0.3"]), 259)

    def test_any_counts_only_the_complete_ipv4_address_space(self):
        count = _count_addresses(["0.0.0.0/0", "::/0"])
        self.assertEqual(count, 2 ** 32)
        self.assertEqual(_excel_safe_count(count), 4294967296)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is optional")
    def test_expanded_rules_contains_address_and_port_count_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); db = Database(root / "state.sqlite"); db.initialize()
            raw = {
                "rule_href": "/rules/1", "ruleset_href": "/rulesets/1",
                "ruleset_name": "APP", "ruleset_scope": "app:APP;env:PRD",
                "src_all_workloads": "true", "dst_iplists": "NETWORKS",
                "services": "All Services",
            }
            db.upsert_rules([raw], "2026-09-12T00:00:00+00:00")
            with db.connect() as connection:
                connection.execute(
                    "INSERT INTO workloads VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                    ("/workloads/1", "host.example", "HOST", "host", "APP", "PRD", "", "", 1,
                     json.dumps(["10.0.0.1", "10.0.0.2"]), "{}", "2026-09-12"),
                )
                connection.execute("INSERT INTO ip_lists VALUES(?,?,?)", ("NETWORKS", "10.1.0.0/24", "2026-09-12"))
            target = generate_workbook(
                db, root, "KEAR", "Application", ["APP"], "PRD", 1,
                date(2026, 9, 12), dangerous_port_lists=["PORTS_TO_CONTROL"],
            )
            from openpyxl import load_workbook
            sheet = load_workbook(target)["Expanded Rules"]
            values = {cell.value: sheet.cell(2, cell.column).value for cell in sheet[1]}
            self.assertEqual(values["nb_src_ips"], 2)
            self.assertEqual(values["nb_dst_ip"], 256)
            self.assertEqual(values["Service Name / Definition"], "0-65535 TCP;0-65535 UDP")
            self.assertEqual(values["nb_ports"], 131072)
            self.assertIn("TCP/22", values["dangerous_ports"])
