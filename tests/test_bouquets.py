import csv
import tempfile
import unittest
from pathlib import Path

from rules_recertify.reporting.bouquets import (
    _factor_side, _labelgroup_dictionary, _workload_dictionary,
    generate_bouquets_workbook,
)


def write_csv(path, fields, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


class BouquetsReportTest(unittest.TestCase):
    def test_dictionaries_have_stable_requested_sort_order(self):
        with tempfile.TemporaryDirectory() as directory:
            workloads = Path(directory) / "workloads.csv"
            write_csv(workloads, ["app", "env", "hostname", "interfaces"], [
                {"app": "B", "env": "DEV", "hostname": "z", "interfaces": "2"},
                {"app": "A", "env": "PRD", "hostname": "b", "interfaces": "1"},
                {"app": "A", "env": "DEV", "hostname": "a", "interfaces": "3"},
            ])
            rows = _workload_dictionary(workloads, {("A", "PRD"), ("A", "DEV"), ("B", "DEV")})
        self.assertEqual([(row["app"], row["env"]) for row in rows],
                         [("A", "DEV"), ("A", "PRD"), ("B", "DEV")])
        groups = [
            {"name": "Z", "key": "role"}, {"name": "B", "key": "app"},
            {"name": "A", "key": "app"},
        ]
        group_rows = _labelgroup_dictionary(groups, {"z": {"R"}, "b": {"R"}, "a": {"R"}})
        self.assertEqual([(row["key"], row["name"]) for row in group_rows],
                         [("app", "A"), ("app", "B"), ("role", "Z")])

    def test_factorized_side_types_each_exact_selector_on_own_line(self):
        rule = {
            "src_all_workloads": "false", "src_labels": "app:APP_A;role:DB.OCP",
            "src_labels_exclusions": "env:DEV", "src_label_groups": "EUR-LGL",
            "src_label_groups_exclusions": "", "src_iplists": "Any (0.0.0.0/0)",
        }
        value, iplists, groups, pairs = _factor_side(
            rule, "src", {"app_a": "app", "db.ocp": "role"}, {"eur-lgl": "loc"},
        )
        self.assertEqual(value.splitlines(), [
            "label.app:APP_A", "label.role:DB.OCP", "label.env:DEV [excluded]",
            "labelgroup.loc:EUR-LGL", "iplist:Any (0.0.0.0/0)",
        ])
        self.assertEqual(iplists, {"Any (0.0.0.0/0)"})
        self.assertEqual(groups, {"EUR-LGL"})
        self.assertEqual(pairs, set())

    def test_workbook_filters_bouquets_and_builds_dictionaries(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl is optional")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            snapshot = root / "raw" / "snapshot"
            snapshot.mkdir(parents=True)
            rule_fields = [
                "ruleset_name", "ruleset_description", "ruleset_enabled", "rule_type",
                "rule_description", "rule_enabled", "src_all_workloads", "src_labels",
                "src_labels_exclusions", "src_label_groups", "src_label_groups_exclusions",
                "src_iplists", "dst_all_workloads", "dst_labels", "dst_labels_exclusions",
                "dst_label_groups", "dst_label_groups_exclusions", "dst_iplists", "services",
            ]
            base = {field: "" for field in rule_fields}
            write_csv(snapshot / "rules_inventory.csv", rule_fields, [
                {**base, "ruleset_name": "BOUQUETS_INFRA", "ruleset_enabled": "true",
                 "rule_type": "sec_rule", "rule_enabled": "true",
                 "src_labels": "app:APP_A;env:PRD", "src_iplists": "NZ3_APP",
                 "dst_label_groups": "EUR-LGL", "services": "SSH"},
                {**base, "ruleset_name": "OTHER", "src_labels": "app:APP_A;env:PRD"},
            ])
            write_csv(snapshot / "labels.csv", ["key", "value"], [
                {"key": "app", "value": "APP_A"}, {"key": "env", "value": "PRD"},
            ])
            group_fields = ["name", "key", "description", "member_labels",
                            "member_label_groups", "fully_expanded_members", "href"]
            write_csv(snapshot / "export_labelgroups.csv", group_fields, [{
                "name": "EUR-LGL", "key": "loc", "description": "Europe",
                "member_labels": "loc:PARIS", "member_label_groups": "",
                "fully_expanded_members": "loc:PARIS", "href": "/groups/1",
            }])
            write_csv(snapshot / "export_iplists.derived.csv", ["name", "include"], [
                {"name": "NZ3_APP", "include": "10.0.0.0/24"},
            ])
            write_csv(snapshot / "export_iplists.csv", ["name", "include"], [
                {"name": "NZ3_APP", "include": "10.0.0.0/24"},
            ])
            write_csv(snapshot / "export_wkld.csv", ["app", "env", "hostname", "interfaces"], [
                {"app": "APP_A", "env": "PRD", "hostname": "host1",
                 "interfaces": "eth0:10.0.0.2"},
            ])
            write_csv(snapshot / "export_services.csv", ["name", "service_ports"], [
                {"name": "SSH", "service_ports": "22 TCP"},
            ])
            target = root / "bouquets.xlsx"
            generate_bouquets_workbook(root / "raw", target)
            workbook = openpyxl.load_workbook(target, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames,
                             ["Bouquets", "iplists", "workloads", "labelgroups", "Metadata"])
            rows = list(workbook["Bouquets"].iter_rows(values_only=True))
            self.assertEqual(len(rows), 2)
            self.assertIn("label.app:APP_A\nlabel.env:PRD\niplist:NZ3_APP", rows[1])
            self.assertEqual(rows[1][-1], "SSH (22 TCP)")
            self.assertEqual(list(workbook["iplists"].iter_rows(values_only=True))[1],
                             ("NZ3_APP", "10.0.0.0/24", "BOUQUETS_INFRA"))
            self.assertEqual(list(workbook["workloads"].iter_rows(values_only=True))[1],
                             ("APP_A", "PRD", "host1", "eth0:10.0.0.2"))
            workbook.close()


if __name__ == "__main__":
    unittest.main()
