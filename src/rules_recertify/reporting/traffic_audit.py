from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Dict, List, Mapping, Sequence, Tuple


SUCCESS_COMPLETE = "SUCCESS"
SUCCESS_WITH_EXCEPTIONS = "SUCCESS_WITH_EXCEPTIONS"
BLOCKING_FAILURE = "WARNING"
SUCCESS_STATUSES = frozenset({SUCCESS_COMPLETE, SUCCESS_WITH_EXCEPTIONS})


def classify_traffic_outcome(
    details: Mapping[str, object], *, allow_empty: bool = False,
) -> Tuple[str, List[str], List[str]]:
    """Return status, blocking reasons, and documented exception reasons."""
    blocking: List[str] = []
    exceptions: List[str] = []

    def count(name: str) -> int:
        return int(details.get(name, 0) or 0)

    if count("total") == 0 and not allow_empty:
        blocking.append("NO_TRAFFIC_RESULTS")
    for field, reason in (
        ("pending", "ASYNC_QUERIES_PENDING"),
        ("expired", "ASYNC_QUERIES_EXPIRED"),
        ("skipped_oversized_ruleset_count", "RULESETS_SKIPPED_OVERSIZED"),
        ("runtime_oversized_ruleset_count", "RULESETS_SKIPPED_RUNTIME_OVERSIZED"),
        ("missing_result_count", "RULES_WITHOUT_RESULT"),
    ):
        if count(field):
            blocking.append(reason)

    for field, reason in (
        ("unknown", "ASYNC_STATUS_UNKNOWN_OR_EMPTY"),
        ("invalid_query_body_count", "INVALID_QUERY_BODY"),
        ("invalid_flows_by_port_count", "INVALID_PORT_DETAILS"),
    ):
        if count(field):
            exceptions.append(reason)

    if blocking:
        return BLOCKING_FAILURE, blocking, exceptions
    if exceptions:
        return SUCCESS_WITH_EXCEPTIONS, blocking, exceptions
    return SUCCESS_COMPLETE, blocking, exceptions


def build_traffic_audit_rows(
    inventory: Sequence[Mapping[str, str]],
    usage_by_rule: Mapping[str, Mapping[str, object]],
    excluded_reasons: Mapping[str, str],
    oversized_rulesets: Mapping[str, str],
    invalid_query_rules: Sequence[str],
    invalid_port_rules: Sequence[str],
) -> Tuple[List[Dict[str, object]], Dict[str, int]]:
    invalid_query = set(invalid_query_rules)
    invalid_ports = set(invalid_port_rules)
    rows: List[Dict[str, object]] = []

    for rule in inventory:
        rule_href = str(rule.get("rule_href", ""))
        ruleset_href = str(rule.get("ruleset_href", ""))
        usage = usage_by_rule.get(rule_href)
        outcome = "PROCESSED"
        reason = "COMPLETED"
        async_status = ""
        flows: object = ""
        batch: object = ""

        if ruleset_href in excluded_reasons:
            outcome = "NOT_IN_SCOPE"
            reason = excluded_reasons[ruleset_href]
        elif ruleset_href in oversized_rulesets:
            outcome = "BLOCKING_PROBLEM"
            reason = oversized_rulesets[ruleset_href]
        elif usage is None:
            outcome = "BLOCKING_PROBLEM"
            reason = "NO_RESULT_RETURNED"
        else:
            async_status = str(usage.get("async_query_status", "")).strip().lower()
            flows = usage.get("flows", "")
            batch = usage.get("_batch", "")
            if rule_href in invalid_query:
                outcome = "DOCUMENTED_EXCEPTION"
                reason = "INVALID_QUERY_BODY"
            elif rule_href in invalid_ports:
                outcome = "DOCUMENTED_EXCEPTION"
                reason = "INVALID_PORT_DETAILS"
            elif async_status == "completed":
                outcome = "PROCESSED"
                reason = "COMPLETED"
            elif async_status in {"", "unknown"}:
                outcome = "DOCUMENTED_EXCEPTION"
                reason = "ASYNC_STATUS_UNKNOWN_OR_EMPTY"
            elif async_status == "expired":
                outcome = "BLOCKING_PROBLEM"
                reason = "ASYNC_QUERY_EXPIRED"
            elif async_status == "pending":
                outcome = "BLOCKING_PROBLEM"
                reason = "ASYNC_QUERY_PENDING"
            else:
                outcome = "DOCUMENTED_EXCEPTION"
                reason = "ASYNC_STATUS_" + async_status.upper()

        rows.append({
            "outcome": outcome,
            "reason": reason,
            "ruleset_name": str(rule.get("ruleset_name", "")),
            "ruleset_scope": str(rule.get("ruleset_scope", "")),
            "ruleset_href": ruleset_href,
            "rule_href": rule_href,
            "rule_description": str(rule.get("rule_description", "")),
            "async_query_status": async_status or ("NOT_APPLICABLE" if usage is None else "EMPTY"),
            "flows": flows,
            "batch": batch,
        })

    counts = Counter(str(row["outcome"]) for row in rows)
    return rows, {
        "processed_rule_count": counts["PROCESSED"],
        "not_in_scope_rule_count": counts["NOT_IN_SCOPE"],
        "documented_exception_rule_count": counts["DOCUMENTED_EXCEPTION"],
        "blocking_rule_count": counts["BLOCKING_PROBLEM"],
        "missing_result_count": sum(
            1 for row in rows if row["reason"] == "NO_RESULT_RETURNED"
        ),
    }


def write_traffic_audit_workbook(
    output: Path, details: Mapping[str, object], rows: Sequence[Mapping[str, object]],
) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to create the traffic audit workbook") from exc

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_rows = [
        ("Run ID", details.get("run_id", "")),
        ("Run type", details.get("run_type", "")),
        ("Final status", details.get("status", "")),
        ("Traffic start", details.get("traffic_start", "")),
        ("Traffic end", details.get("traffic_end", "")),
        ("Total async results", details.get("total", 0)),
        ("Completed", details.get("completed", 0)),
        ("Pending", details.get("pending", 0)),
        ("Expired", details.get("expired", 0)),
        ("Unknown / empty", details.get("unknown", 0)),
        ("Invalid query bodies", details.get("invalid_query_body_count", 0)),
        ("Invalid port details", details.get("invalid_flows_by_port_count", 0)),
        ("Blocking reasons", ";".join(details.get("blocking_reasons", []))),
        ("Documented exceptions", ";".join(details.get("documented_exception_reasons", []))),
    ]
    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    for item in summary_rows:
        summary.append(list(item))
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 80

    columns = list(rows[0]) if rows else ["outcome", "reason", "rule_href"]
    sheets = (
        ("Rules Treated", [row for row in rows if row.get("outcome") == "PROCESSED"]),
        ("Not In Scope", [row for row in rows if row.get("outcome") == "NOT_IN_SCOPE"]),
        ("Problematic Rules", [row for row in rows if row.get("outcome") in {"DOCUMENTED_EXCEPTION", "BLOCKING_PROBLEM"}]),
        ("All Rules", list(rows)),
    )
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for name, sheet_rows in sheets:
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for row in sheet_rows:
            sheet.append([row.get(column, "") for column in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(columns, 1):
            maximum = max([len(column)] + [len(str(row.get(column, ""))) for row in sheet_rows])
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(maximum + 2, 80)

    workbook.save(output)
    return output
