"""Generate the infrastructure-bouquet ruleset inventory workbook."""
from __future__ import annotations

import importlib.util
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Sequence, Set, Tuple

from ..workloader.csvio import parse_bool, read_rows
from .rule_search import _write_sheet


RULE_FIELDS = (
    "ruleset_name", "ruleset_description", "ruleset_enabled", "rule_type",
    "rule_description", "rule_enabled", "src_all_workloads", "src_labels",
    "src_labels_exclusions", "src_label_groups", "src_label_groups_exclusions",
    "src_iplists", "dst_all_workloads", "dst_labels", "dst_labels_exclusions",
    "dst_label_groups", "dst_label_groups_exclusions", "dst_iplists", "services",
)
LABEL_GROUP_FIELDS = (
    "name", "key", "description", "member_labels", "member_label_groups",
    "fully_expanded_members", "href",
)
WORKLOAD_FIELDS = ("app", "env", "hostname", "interfaces")
IPLIST_FIELDS = ("name", "include")


def generate_bouquets_workbook(raw_dir: Path, output: Path) -> Path:
    """Export BOUQUETS_* rules and the references used to interpret them."""
    if importlib.util.find_spec("openpyxl") is None:
        raise RuntimeError("openpyxl is required for bouquet report generation")
    snapshot = raw_dir / "snapshot"
    paths = {
        "rules": snapshot / "rules_inventory.csv",
        "labels": snapshot / "labels.csv",
        "groups": snapshot / "export_labelgroups.csv",
        "iplists": snapshot / "export_iplists.derived.csv",
        "iplists_full": snapshot / "export_iplists.csv",
        "workloads": snapshot / "export_wkld.csv",
    }
    missing = [str(path) for path in paths.values() if not path.is_file() or not path.stat().st_size]
    if missing:
        raise ValueError("bouquet snapshot file is missing or empty: " + ", ".join(missing))

    rules = [
        row for row in read_rows(paths["rules"], RULE_FIELDS)
        if row["ruleset_name"].strip().upper().startswith("BOUQUETS_")
    ]
    labels = list(read_rows(paths["labels"], ("key", "value")))
    groups = list(read_rows(paths["groups"], LABEL_GROUP_FIELDS))
    label_types = _type_index(labels)
    group_types = _type_index(groups)

    bouquet_rows: List[Dict[str, object]] = []
    used_iplists: Dict[str, Set[str]] = {}
    used_groups: Dict[str, Set[str]] = {}
    workload_pairs: Set[Tuple[str, str]] = set()
    for rule in rules:
        source, source_iplists, source_groups, source_pairs = _factor_side(
            rule, "src", label_types, group_types,
        )
        destination, destination_iplists, destination_groups, destination_pairs = _factor_side(
            rule, "dst", label_types, group_types,
        )
        ruleset = rule["ruleset_name"]
        for name in source_iplists | destination_iplists:
            used_iplists.setdefault(name.casefold(), set()).add(ruleset)
        for name in source_groups | destination_groups:
            used_groups.setdefault(name.casefold(), set()).add(ruleset)
        workload_pairs.update(source_pairs | destination_pairs)
        bouquet_rows.append({
            "ruleset_name": ruleset,
            "ruleset_description": rule["ruleset_description"],
            "ruleset_enabled": rule["ruleset_enabled"],
            "rule_type": rule["rule_type"],
            "rule_description": rule["rule_description"],
            "rule_enabled": rule["rule_enabled"],
            "Source": source,
            "Destination": destination,
            "services": rule["services"],
        })

    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "Bouquets", bouquet_rows)
    _write_sheet(workbook, "iplists", _iplist_dictionary(
        (paths["iplists"], paths["iplists_full"]), used_iplists,
    ))
    _write_sheet(workbook, "workloads", _workload_dictionary(paths["workloads"], workload_pairs))
    _write_sheet(workbook, "labelgroups", _labelgroup_dictionary(groups, used_groups))
    _write_sheet(workbook, "Metadata", [{
        "rules inventory": str(paths["rules"]), "bouquet prefix": "BOUQUETS_",
        "matching rules": len(rules), "generated at UTC": datetime.now(timezone.utc).isoformat(),
    }])
    from openpyxl.styles import Alignment
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row_number in range(2, sheet.max_row + 1):
            line_count = max((str(sheet.cell(row_number, column).value or "").count("\n") + 1
                              for column in range(1, sheet.max_column + 1)), default=1)
            sheet.row_dimensions[row_number].height = min(15 * line_count, 120)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    workbook.save(temporary)
    temporary.replace(output)
    return output


def _type_index(rows: Iterable[Mapping[str, str]]) -> Dict[str, str]:
    output: Dict[str, str] = {}
    for row in rows:
        name = str(row.get("value", row.get("name", ""))).strip()
        key = str(row.get("key", "")).strip().casefold()
        if name and key:
            output[name.casefold()] = key
    return output


def _parts(value: object) -> List[str]:
    return [part.strip() for part in re.split(r"[;\n]", str(value or "")) if part.strip()]


def _clean_reference(value: str) -> str:
    value = re.sub(r"^(?:IP\s*List|Label\s*Group)\s*:\s*", "", value, flags=re.I)
    return re.sub(r"\s*\(/[^)]*\)\s*$", "", value).strip()


def _typed_label(value: str, types: Mapping[str, str]) -> Tuple[str, str, str]:
    explicit_key, separator, explicit_value = value.partition(":")
    name = explicit_value.strip() if separator else _clean_reference(value)
    key = (explicit_key.strip() if separator else types.get(name.casefold(), "unknown")).casefold()
    return f"label.{key}:{name}", key, name


def _factor_side(
    rule: Mapping[str, str], prefix: str, label_types: Mapping[str, str],
    group_types: Mapping[str, str],
) -> Tuple[str, Set[str], Set[str], Set[Tuple[str, str]]]:
    output: List[str] = []
    iplists: Set[str] = set()
    groups: Set[str] = set()
    apps: Set[str] = set()
    envs: Set[str] = set()
    if parse_bool(rule.get(f"{prefix}_all_workloads", "")):
        output.append("workload:All workloads")
    for field, exclusion in ((f"{prefix}_labels", False), (f"{prefix}_labels_exclusions", True)):
        for value in _parts(rule.get(field, "")):
            display, key, name = _typed_label(value, label_types)
            output.append(display + (" [excluded]" if exclusion else ""))
            if not exclusion and key == "app":
                apps.add(name)
            elif not exclusion and key == "env":
                envs.add(name)
    for field, exclusion in ((f"{prefix}_label_groups", False), (f"{prefix}_label_groups_exclusions", True)):
        for value in _parts(rule.get(field, "")):
            name = _clean_reference(value)
            key = group_types.get(name.casefold(), "unknown")
            output.append(f"labelgroup.{key}:{name}" + (" [excluded]" if exclusion else ""))
            groups.add(name)
    for value in _parts(rule.get(f"{prefix}_iplists", "")):
        name = _clean_reference(value)
        output.append(f"iplist:{name}")
        iplists.add(name)
    return ("\n".join(dict.fromkeys(output)), iplists, groups,
            {(app, env) for app in apps for env in envs})


def _iplist_dictionary(
    paths: Sequence[Path], used: Mapping[str, Set[str]],
) -> List[Dict[str, str]]:
    rows = []
    emitted: Set[str] = set()
    for path in paths:
        for row in read_rows(path, IPLIST_FIELDS):
            identity = row["name"].casefold()
            rulesets = used.get(identity)
            if rulesets and identity not in emitted:
                rows.append({"name": row["name"], "include": row["include"],
                             "used in": "\n".join(sorted(rulesets))})
                emitted.add(identity)
    return rows


def _workload_dictionary(path: Path, pairs: Set[Tuple[str, str]]) -> List[Dict[str, str]]:
    wanted = {(app.casefold(), env.casefold()) for app, env in pairs}
    return [{field: row[field] for field in WORKLOAD_FIELDS}
            for row in read_rows(path, WORKLOAD_FIELDS)
            if (row["app"].casefold(), row["env"].casefold()) in wanted]


def _labelgroup_dictionary(
    groups: Sequence[Mapping[str, str]], used: Mapping[str, Set[str]],
) -> List[Dict[str, str]]:
    rows = []
    for group in groups:
        rulesets = used.get(str(group.get("name", "")).casefold())
        if rulesets:
            rows.append({
                "name": str(group.get("name", "")), "key": str(group.get("key", "")),
                "description": str(group.get("description", "")),
                "member_labels": str(group.get("member_labels", "")),
                "member_label_groups": str(group.get("member_label_groups", "")),
                "fully_expanded_members": str(group.get("fully_expanded_members", "")),
                "used in": "\n".join(sorted(rulesets)),
            })
    return rows
