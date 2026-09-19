"""Bulk workbook generation driven by a Microcosmos XLSX export."""
from __future__ import annotations

import importlib.util
import json
import logging
import re
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from ..history.database import Database
from ..workloader.csvio import read_rows
from .workbook import ReportingDependencyError, _rule_matches, generate_workbook

REQUIRED_COLUMNS = (
    "Kear Id", "Application Name", "Module", "Account", "Account Leader",
    "Microsegmentation Solution", "Environment", "Entity",
)
LOG = logging.getLogger(__name__)


def generate_microcosmos_reports(
    db: Database,
    source: Path,
    output_dir: Path,
    raw_dir: Path,
    lookback_days: int,
    as_of: date,
    timestamp: Optional[str] = None,
    dangerous_port_lists: Sequence[str] = (),
    device: str = "",
    permissive_rule_max_ips: int = 255,
) -> List[Path]:
    """Generate one PRD/NONPRD report per Entity and non-empty KEAR ID."""
    rows = _read_microcosmos(source)
    labels = _known_application_labels(db, raw_dir)
    rule_rows = _rule_rows(db)
    statuses: Dict[int, str] = {}
    groups: Dict[Tuple[str, str, str], List[Mapping[str, str]]] = defaultdict(list)
    for fallback_number, row in enumerate(rows, 2):
        row_number = int(row.get("__row_number__", fallback_number))
        kear = row["Kear Id"].strip()
        if not kear:
            statuses[row_number] = "SKIPPED: empty Kear Id"
            continue
        category = "PRD" if row["Environment"].strip().upper() == "PRD" else "NONPRD"
        environment = row["Environment"].strip()
        if not environment:
            statuses[row_number] = "SKIPPED: empty Environment"
            continue
        matches = _labels_for_module(row["Module"], labels)
        if not matches:
            statuses[row_number] = f"SKIPPED: no application label matches module {row['Module']!r}"
            continue
        matching_labels = [
            label for label in matches
            if any(_rule_matches(rule, [(label, environment)]) for rule in rule_rows)
        ]
        if not matching_labels:
            statuses[row_number] = "SKIPPED: no matching ruleset/rule in SQLite"
            continue
        prepared = dict(row)
        prepared["__labels__"] = "\n".join(matching_labels)
        prepared["__row_number__"] = str(row_number)
        groups[(row["Entity"].strip() or "UNSPECIFIED", kear, category)].append(prepared)

    root = output_dir / (timestamp or datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"))
    generated: List[Path] = []
    for (entity, kear, category), group in sorted(groups.items()):
        names = {row["Application Name"].strip() for row in group if row["Application Name"].strip()}
        if len(names) != 1:
            for row in group:
                statuses[int(row["__row_number__"])] = "SKIPPED: inconsistent or empty Application Name"
            continue
        pairs: List[Tuple[str, str]] = []
        for row in group:
            environment = row["Environment"].strip()
            pairs.extend((label, environment) for label in row["__labels__"].splitlines())
        pairs = list(dict.fromkeys(pairs))
        target_dir = root / category / _safe_component(entity)
        try:
            target = generate_workbook(
                db, target_dir, kear, next(iter(names)),
                [pair[0] for pair in pairs], [pair[1] for pair in pairs],
                lookback_days, as_of, raw_dir=raw_dir, filename_environment=category,
                dangerous_port_lists=dangerous_port_lists,
                device=device, permissive_rule_max_ips=permissive_rule_max_ips,
            )
        except (ValueError, RuntimeError) as exc:
            LOG.error("Skipping Microcosmos report %s/%s/%s: %s", entity, kear, category, exc)
            for row in group:
                statuses[int(row["__row_number__"])] = f"ERROR: {exc}"
            continue
        generated.append(target)
        for row in group:
            statuses[int(row["__row_number__"])] = f"PROCESSED: {target.relative_to(root)}"
    audit = _write_audit_workbook(source, root, statuses)
    LOG.info("Microcosmos batch completed: %s reports; audit=%s", len(generated), audit)
    return [audit, *generated]


def _read_microcosmos(source: Path) -> List[Dict[str, str]]:
    if importlib.util.find_spec("openpyxl") is None:
        raise ReportingDependencyError("openpyxl is required to read the Microcosmos XLSX export")
    if not source.is_file():
        raise ValueError(f"Microcosmos workbook not found: {source}")
    from openpyxl import load_workbook
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(iterator, ())]
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"{source}: missing columns: {', '.join(missing)}")
    indexes = {column: headers.index(column) for column in REQUIRED_COLUMNS}
    return [
        {**{column: str(values[index] or "").strip() if index < len(values) else ""
             for column, index in indexes.items()}, "__row_number__": str(number)}
        for number, values in enumerate(iterator, 2)
        if any(value is not None and str(value).strip() for value in values)
    ]


def _known_application_labels(db: Database, raw_dir: Optional[Path] = None) -> List[str]:
    labels = set()
    with db.connect() as connection:
        labels.update(str(row[0]).strip() for row in connection.execute("SELECT DISTINCT app FROM workloads"))
        for row in connection.execute("SELECT raw_json FROM rules WHERE is_present=1"):
            raw = json.loads(row[0])
            for value in _strings(raw):
                labels.update(match.strip() for match in re.findall(r"(?:^|[;\n])app:([^;\n]+)", value, re.I))
    if raw_dir:
        timestamped = sorted(
            (path for path in raw_dir.glob("*/labels.csv")
             if re.fullmatch(r"\d{8}T\d{6}Z-[0-9A-Fa-f]{8}", path.parent.name) and path.is_file()),
            key=lambda path: path.parent.name, reverse=True,
        )
        snapshot = raw_dir / "snapshot" / "labels.csv"
        label_exports = ([snapshot] if snapshot.is_file() and snapshot.stat().st_size else []) + timestamped
        if label_exports:
            labels.update(
                row["value"] for row in read_rows(label_exports[0], ("key", "value"))
                if row["key"].strip().casefold() == "app" and row["value"].strip()
            )
    return sorted((label for label in labels if label), key=str.casefold)


def _rule_rows(db: Database) -> List[Dict[str, str]]:
    with db.connect() as connection:
        return [{"raw_json": str(row[0])} for row in connection.execute(
            "SELECT raw_json FROM rules WHERE is_present=1"
        )]


def _write_audit_workbook(source: Path, root: Path, statuses: Mapping[int, str]) -> Path:
    from openpyxl import load_workbook
    workbook = load_workbook(source)
    sheet = workbook.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    title = "Rules Recertify Status"
    column = headers.index(title) + 1 if title in headers else len(headers) + 1
    sheet.cell(1, column, title)
    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row_number, column, statuses.get(row_number, "SKIPPED: empty row"))
    root.mkdir(parents=True, exist_ok=True)
    target = root / f"{source.stem}.rules-recertify-status.xlsx"
    temporary = target.with_suffix(".xlsx.tmp")
    workbook.save(temporary)
    temporary.replace(target)
    return target


def _strings(value: object) -> Iterable[str]:
    if isinstance(value, str):
        yield value
    elif isinstance(value, Mapping):
        for nested in value.values():
            yield from _strings(nested)
    elif isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        for nested in value:
            yield from _strings(nested)


def _labels_for_module(module: str, labels: Sequence[str]) -> List[str]:
    wanted = module.strip().casefold()
    if not wanted:
        return []
    return [label for label in labels if len(label.split("_", 2)) == 3 and label.split("_", 2)[2].casefold() == wanted]


def _safe_component(value: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("._")
    return safe or "UNSPECIFIED"
