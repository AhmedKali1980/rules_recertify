"""Generate rule-compliance workbooks from the materialized policy snapshot."""
from __future__ import annotations

import csv
import importlib.util
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Mapping, Optional, Sequence, Tuple

from ..workloader.csvio import parse_bool, read_rows
from .dangerous_ports import dangerous_ports
from .rule_search import _write_sheet
from .workbook import _expand_services, _load_report_services


RULE_FIELDS = (
    "ruleset_name", "ruleset_scope", "ruleset_enabled", "rule_type", "rule_enabled",
    "unscoped_consumers", "src_all_workloads", "src_labels", "src_iplists",
    "dst_all_workloads", "dst_labels", "dst_iplists", "services",
)
LABEL_GROUP_FIELDS = (
    "name", "key", "description", "member_labels", "member_label_groups",
    "fully_expanded_members", "href",
)
MICROCOSMOS_COLUMNS = (
    "Kear Id", "Application Name", "Module", "Account", "Account Leader",
    "Microsegmentation Solution", "Environment", "Entity", "Program (MIO)",
)
MICROCOSMOS_REQUIRED_COLUMNS = MICROCOSMOS_COLUMNS[:-1]


def generate_compliance_workbook(
    raw_dir: Path, output: Path, dangerous_port_lists: Sequence[str],
    input_path: Optional[Path] = None,
) -> Path:
    """Evaluate matching snapshot rules for optional app/environment input rows."""
    if importlib.util.find_spec("openpyxl") is None:
        raise RuntimeError("openpyxl is required for compliance report generation")
    snapshot = raw_dir / "snapshot"
    required = [snapshot / "rules_inventory.csv", snapshot / "export_labelgroups.csv"]
    if input_path is None:
        required.append(snapshot / "labels.csv")
    missing = [str(path) for path in required if not path.is_file() or not path.stat().st_size]
    if missing:
        raise ValueError("compliance snapshot file is missing or empty: " + ", ".join(missing))
    input_rows, input_columns = _load_scope_input(input_path, snapshot / "labels.csv")
    rules = list(read_rows(snapshot / "rules_inventory.csv", RULE_FIELDS))
    groups = list(read_rows(snapshot / "export_labelgroups.csv", LABEL_GROUP_FIELDS))
    services, service_reference = _load_report_services(raw_dir)
    intra_zone_rulesets = {
        _ruleset_identity(row) for row in rules
        if _has_intra_zone_iplist(row.get("src_iplists", ""))
        or _has_intra_zone_iplist(row.get("dst_iplists", ""))
    }

    rows: List[Dict[str, object]] = []
    for input_row in input_rows:
        app = _get_case_insensitive(input_row, "app").strip()
        env = str(input_row.get("__compliance_env", _get_case_insensitive(input_row, "env"))).strip()
        programs = _programs_for_app(app, groups)
        matched = [row for row in rules if _rule_matches_scope(row, app, env)]
        primary_columns = (
            [*MICROCOSMOS_COLUMNS, "app"]
            if "Program (MIO)" in input_columns else input_columns
        )
        input_values = {column: input_row.get(column, "") for column in primary_columns}
        extra_values = {
            column: input_row.get(column, "")
            for column in input_columns if column not in primary_columns
        }
        if not matched:
            rows.append({
                **input_values, "Program (KPI)": programs,
                "Match status": "APP_NOT_FOUND" if not app else "NO_MATCH",
                "Warning": "",
                "Global Consumer in use": "FALSE",
                **{field: "" for field in RULE_FIELDS[:5]},
                "Source": "", "Destination": "", "Source contains User": "FALSE",
                "Destination contains User": "FALSE",
                "services": "",
                **{catalog: "" for catalog in dangerous_port_lists},
                "Protection type": "",
                **extra_values,
            })
            continue
        for rule in matched:
            source, source_kind = _factor_side(rule, "src")
            destination, destination_kind = _factor_side(rule, "dst")
            expanded_services = _expand_services(rule["services"], services)
            rows.append({
                **input_values,
                "Program (KPI)": programs,
                "Match status": "MATCHED",
                "Warning": _extrascope_warning(rule),
                "Global Consumer in use": str(_global_consumer_in_use(rule)).upper(),
                "ruleset_name": rule["ruleset_name"],
                "ruleset_scope": rule["ruleset_scope"],
                "ruleset_enabled": rule["ruleset_enabled"],
                "rule_type": rule["rule_type"],
                "rule_enabled": rule["rule_enabled"],
                "Source": source,
                "Destination": destination,
                "Source contains User": str(
                    source_kind == "iplist" and _contains_user_iplist(source)
                ).upper(),
                "Destination contains User": str(
                    destination_kind == "iplist" and _contains_user_iplist(destination)
                ).upper(),
                "services": rule["services"],
                **{
                    catalog: dangerous_ports(expanded_services, [catalog])
                    for catalog in dangerous_port_lists
                },
                "Protection type": (
                    "Intra-zone" if _ruleset_identity(rule) in intra_zone_rulesets else ""
                ),
                **extra_values,
            })

    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "Compliance", rows)
    _write_sheet(workbook, "Metadata", [{
        "rules inventory": str(snapshot / "rules_inventory.csv"),
        "input": str(input_path) if input_path else str(snapshot / "labels.csv"),
        "service reference": service_reference,
        "dangerous port lists": ";".join(dangerous_port_lists),
        "generated at UTC": datetime.now(timezone.utc).isoformat(),
    }])
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    workbook.save(temporary)
    temporary.replace(output)
    return output


def _load_scope_input(input_path: Optional[Path], labels_path: Path) -> Tuple[List[Dict[str, str]], List[str]]:
    if input_path is None:
        rows = [
            {"app": row["value"]}
            for row in read_rows(labels_path, ("key", "value"))
            if row["key"].strip().casefold() == "app" and row["value"].strip()
        ]
        return _deduplicate_inputs(rows, ("app",)), ["app"]
    if not input_path.is_file():
        raise ValueError(f"compliance input file not found: {input_path}")
    if input_path.suffix.casefold() == ".xlsx":
        if importlib.util.find_spec("openpyxl") is None:
            raise RuntimeError("openpyxl is required to read an XLSX compliance input")
        from openpyxl import load_workbook
        workbook = load_workbook(input_path, read_only=True, data_only=True)
        iterator = workbook.active.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator, ())]
        values = list(iterator)
        workbook.close()
        rows = [
            {header: str(row[index] or "").strip() if index < len(row) else ""
             for index, header in enumerate(headers)}
            for row in values if any(str(value or "").strip() for value in row)
        ]
    elif input_path.suffix.casefold() == ".csv":
        with input_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            raw_headers = list(reader.fieldnames or [])
            headers = [str(value or "").strip() for value in raw_headers]
            rows = [
                {
                    header: str(row.get(raw_header, "") or "").strip()
                    for raw_header, header in zip(raw_headers, headers)
                }
                for row in reader
            ]
    else:
        raise ValueError("compliance input must be a CSV or XLSX file")
    if any(not header for header in headers):
        raise ValueError("compliance input contains an empty column header")
    if len({header.casefold() for header in headers}) != len(headers):
        raise ValueError("compliance input contains duplicate column headers")
    header_names = {header.casefold(): header for header in headers}
    if input_path.suffix.casefold() == ".xlsx" and {
        name.casefold() for name in MICROCOSMOS_REQUIRED_COLUMNS
    }.issubset(header_names):
        return _prepare_microcosmos_input(rows, headers, labels_path)
    if "app" not in header_names:
        raise ValueError(
            "compliance input must contain an app column, or be a Microcosmos XLSX "
            "with Kear Id, Application Name, Module, Account, Account Leader, "
            "Microsegmentation Solution, Environment, and Entity columns"
        )
    rows = [row for row in rows if _get_case_insensitive(row, "app").strip()]
    if not rows:
        raise ValueError("compliance input contains no non-empty app value")
    return rows, headers


def _prepare_microcosmos_input(
    rows: Sequence[Mapping[str, str]], headers: Sequence[str], labels_path: Path,
) -> Tuple[List[Dict[str, str]], List[str]]:
    """Translate Microcosmos Module/Environment rows into compliance app/env scopes."""
    if not labels_path.is_file() or not labels_path.stat().st_size:
        raise ValueError(f"compliance snapshot file is missing or empty: {labels_path}")
    labels = list(read_rows(labels_path, ("key", "value")))
    apps_by_module: Dict[str, List[str]] = {}
    environments: Dict[str, str] = {}
    for label in labels:
        key, value = label["key"].strip().casefold(), label["value"].strip()
        if not value:
            continue
        if key == "app":
            module = _module_from_app(value)
            if module:
                apps_by_module.setdefault(module.casefold(), []).append(value)
        elif key == "env":
            environments.setdefault(value.casefold(), value)

    canonical_by_name = {column.casefold(): column for column in MICROCOSMOS_COLUMNS}
    source_by_name = {header.casefold(): header for header in headers}
    extras = [
        header for header in headers
        if header.casefold() not in canonical_by_name and header.casefold() != "program"
    ]
    output_columns = [*MICROCOSMOS_COLUMNS, "app", *extras]
    prepared: List[Dict[str, str]] = []
    for source in rows:
        module = _get_case_insensitive(source, "Module").strip()
        environment = _get_case_insensitive(source, "Environment").strip()
        apps = list(dict.fromkeys(apps_by_module.get(module.casefold(), []))) or [""]
        for app in apps:
            output: Dict[str, str] = {}
            for column in MICROCOSMOS_COLUMNS:
                source_name = source_by_name.get(column.casefold())
                if column == "Program (MIO)" and source_name is None:
                    source_name = source_by_name.get("program")
                output[column] = str(source.get(source_name, "") or "") if source_name else ""
            for column in extras:
                output[column] = str(source.get(column, "") or "")
            output["app"] = app
            output["__compliance_env"] = environments.get(environment.casefold(), environment)
            prepared.append(output)
    if not prepared:
        raise ValueError("compliance input contains no non-empty Microcosmos row")
    return prepared, output_columns


def _module_from_app(app: str) -> str:
    """Return the part after the second underscore in an application label."""
    components = app.split("_", 2)
    return components[2].strip() if len(components) == 3 else ""


def _deduplicate_inputs(rows: Sequence[Dict[str, str]], keys: Sequence[str]) -> List[Dict[str, str]]:
    output: List[Dict[str, str]] = []
    seen = set()
    for row in rows:
        identity = tuple(row.get(key, "").casefold() for key in keys)
        if identity not in seen:
            output.append(row)
            seen.add(identity)
    return output


def _get_case_insensitive(row: Mapping[str, object], name: str) -> str:
    wanted = name.casefold()
    return next((str(value or "") for key, value in row.items() if key.casefold() == wanted), "")


def _labels(value: object) -> Dict[str, set[str]]:
    output: Dict[str, set[str]] = {}
    for component in re.split(r"[;\n]", str(value or "")):
        key, separator, item = component.strip().partition(":")
        if separator and key.strip() and item.strip():
            output.setdefault(key.strip().casefold(), set()).add(item.strip().casefold())
    return output


def _rule_matches_scope(rule: Mapping[str, str], app: str, env: str) -> bool:
    wanted_app, wanted_env = app.casefold(), env.casefold()
    for field in ("ruleset_scope", "src_labels", "dst_labels"):
        labels = _labels(rule.get(field, ""))
        if wanted_app not in labels.get("app", set()):
            continue
        if not wanted_env or wanted_env in labels.get("env", set()):
            return True
    return False


def _programs_for_app(app: str, groups: Sequence[Mapping[str, str]]) -> str:
    programs = []
    for group in groups:
        name = str(group.get("name", "")).strip()
        if not name.upper().startswith("KPI_PROGRAM_"):
            continue
        members = ";".join((
            str(group.get("member_labels", "")),
            str(group.get("fully_expanded_members", "")),
        ))
        if re.search(rf"(?<![A-Za-z0-9_]){re.escape(app)}(?![A-Za-z0-9_])", members, re.I):
            programs.append(name)
    return ";".join(dict.fromkeys(programs))


def _factor_side(rule: Mapping[str, str], prefix: str) -> Tuple[str, str]:
    if parse_bool(rule.get(f"{prefix}_all_workloads", "")):
        context = _ruleset_scope_context(rule.get("ruleset_scope", ""))
        return (f"All workloads ({context})" if context else "All workloads"), "all_workloads"
    if rule.get(f"{prefix}_iplists", "").strip():
        return rule[f"{prefix}_iplists"], "iplist"
    if rule.get(f"{prefix}_labels", "").strip():
        return rule[f"{prefix}_labels"], "labels"
    return "", ""


def _ruleset_scope_context(scope: object) -> str:
    """Render app/env labels from a ruleset scope in a stable readable order."""
    values: Dict[str, List[str]] = {"app": [], "env": []}
    for component in re.split(r"[;\n]", str(scope or "")):
        key, separator, value = component.strip().partition(":")
        normalized_key = key.strip().casefold()
        value = value.strip()
        if separator and normalized_key in values and value and value not in values[normalized_key]:
            values[normalized_key].append(value)
    return "; ".join(
        f"{key}:{', '.join(values[key])}" for key in ("app", "env") if values[key]
    )


def _extrascope_warning(rule: Mapping[str, str]) -> str:
    """Explain the critical Global Consumer condition when it is detected."""
    if _global_consumer_in_use(rule):
        return "WARNING: Global Consumer applies the rule to all managed workloads in the PCE"
    return ""


def _global_consumer_in_use(rule: Mapping[str, str]) -> bool:
    """Return whether an Extrascope rule consumes every managed PCE workload."""
    return bool(
        parse_bool(rule.get("unscoped_consumers", ""))
        and parse_bool(rule.get("src_all_workloads", ""))
    )


def _iplist_names(value: str) -> List[str]:
    names = []
    for part in re.split(r"[;\n]", value):
        name = part.strip().split(" (", 1)[0].strip()
        if name.casefold().startswith("ip list:"):
            name = name.partition(":")[2].strip()
        if name:
            names.append(name)
    return names


def _contains_user_iplist(value: str) -> bool:
    return any(name.upper().startswith(("U_", "USR_")) for name in _iplist_names(value))


def _has_intra_zone_iplist(value: str) -> bool:
    return any(name.upper().startswith(("NZ3_", "ZNOT_NZ3")) for name in _iplist_names(value))


def _ruleset_identity(rule: Mapping[str, str]) -> str:
    return rule.get("ruleset_href", "").strip() or rule.get("ruleset_name", "").strip()
